from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parent
OUTPUTS_DIR = PROJECT_DIR / "outputs"
PLOTS_DIR = PROJECT_DIR / "plots"
REPORT_PATH = PROJECT_DIR / "report_lab4_map_saved_models.docx"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def fmt_int(value: float | int) -> str:
    return f"{int(round(value)):,}".replace(",", " ")


def fmt_money(value: float | int) -> str:
    return f"{value / 1_000_000:,.2f}".replace(",", " ") + " млн"


def fmt_pct(value: float | int) -> str:
    return f"{float(value):.2f}%"


def fmt_metric(value: float | int, digits: int = 3) -> str:
    return f"{float(value):.{digits}f}"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_table(table, header_fill: str = "D9EAF7") -> None:
    table.style = "Table Grid"
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(color)
    run_pr.append(underline)
    run.append(run_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_paragraph(doc: Document, text: str = "", bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        first.font.name = "Times New Roman"
        first.font.size = Pt(11)
        rest = paragraph.add_run(text[len(bold_prefix):])
        rest.font.name = "Times New Roman"
        rest.font.size = Pt(11)
    else:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(31, 78, 121)
        run.font.bold = True
        run.font.size = Pt(16 if level == 1 else 13)
    return heading


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.name = "Times New Roman"


def add_picture(doc: Document, filename: str, caption: str, width: float = 6.25) -> None:
    path = PLOTS_DIR / filename
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(f"- {item}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(f"{index}. {item}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for i, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0 if i < len(lines) - 1 else 8)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)


def add_passport_table(doc: Document, stats: dict) -> None:
    rows = [
        ("1", "Название датасета", "Retail RTO Forecast: история РТО по магазинам сети"),
        ("2", "Источник", "Локальный файл проекта data/train.csv"),
        ("3", "Количество наблюдений", f"{fmt_int(stats['raw_rows'])} строк в исходном наборе; {fmt_int(stats['processed_rows'])} строк после лагов"),
        ("4", "Количество признаков", f"{stats['raw_cols']} столбцов исходно; {stats['processed_cols']} столбцов после подготовки"),
        ("5", "Тип задачи", "Регрессия, прогноз непрерывного значения РТО"),
        ("6", "Целевая переменная", "РТО, числовой показатель месячного оборота магазина"),
        ("7", "Доля целевого класса", "Не применяется, так как задача не классификационная"),
        ("8", "Категориальные / числовые признаки", f"{stats['categorical_feature_count_raw']} категориальных и {stats['numeric_feature_count_raw']} числовых признаков в исходных данных"),
        ("9", "Пропуски и выбросы", f"Пропусков: {stats['missing_raw']}; IQR-выбросов по РТО: {fmt_int(stats['target_outliers_iqr'])} ({fmt_pct(stats['target_outliers_iqr_share'])})"),
        ("10", "Предварительный вывод", "Данные полные, но РТО имеет правостороннее распределение; выбросы отражают крупные магазины и сохранены"),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ["№", "Характеристика набора данных", "Значение / описание"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    format_table(table)


def add_preprocessing_table(doc: Document) -> None:
    rows = [
        ("1", "Обработка пропусков", "Проверка isna; строк с пропусками нет", "Пропуски отсутствуют в исходном и обработанном наборе"),
        ("2", "Обработка выбросов", "IQR-анализ целевой переменной; значения сохранены", "Крупные РТО являются бизнес-реальностью сети и важны для прогноза"),
        ("3", "Кодирование категорий", "Ordinal encoding для возраста/площади; one-hot encoding для региона", "Категории переведены в числовой вид, регион не навязывает искусственный порядок"),
        ("4", "Масштабирование", "StandardScaler для числовых признаков из сохраненного scaler.pkl", "Нужно для линейной регрессии и MLP; лес использует исходный масштаб"),
        ("5", "Разделение выборок", "Holdout 80/20, random_state=42", "Единая тестовая часть позволяет сравнить модели корректно"),
        ("6", "Балансировка классов", "Не выполнялась", "В задаче регрессии нет целевых классов"),
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Шаг", "Действие", "Инструменты / метод", "Обоснование"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    format_table(table)


def add_models_table(doc: Document) -> None:
    rows = [
        ("LinearRegression", "sklearn.linear_model.LinearRegression", "fit_intercept=True; признаки масштабированы", "Базовая интерпретируемая модель регрессии"),
        ("RandomForestRegressor", "sklearn.ensemble.RandomForestRegressor", "n_estimators=100; max_depth=20; min_samples_split=4; random_state=42", "Нелинейная табличная модель и источник feature importance"),
        ("TorchMLP", "PyTorch MLP", "128 -> 64 -> 1; BatchNorm; Dropout; target_transform=log1p", "Лучшая модель по MAPE; использована для итогового прогноза"),
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Модель", "Библиотека", "Гиперпараметры / артефакт", "Назначение"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    format_table(table)


def add_metrics_table(doc: Document, metrics: pd.DataFrame) -> None:
    ordered = metrics.sort_values("MAPE").reset_index(drop=True)
    table = doc.add_table(rows=1, cols=6)
    headers = ["Модель", "MAE", "RMSE", "R²", "MAPE", "Вывод"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for _, row in ordered.iterrows():
        conclusion = "Лучшая по MAPE" if row["model"] == ordered.iloc[0]["model"] else "Качество близкое" if row["MAPE"] < 5 else "Базовый ориентир"
        values = [
            row["model"],
            fmt_money(row["MAE"]),
            fmt_money(row["RMSE"]),
            fmt_metric(row["R2"], 4),
            fmt_pct(row["MAPE"]),
            conclusion,
        ]
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value)
    best_cells = table.add_row().cells
    best_values = ["Лучшая модель", ordered.iloc[0]["model"], "", "", fmt_pct(ordered.iloc[0]["MAPE"]), "Выбрана для test.csv"]
    for i, value in enumerate(best_values):
        set_cell_text(best_cells[i], value, bold=i in (0, 1, 4, 5))
    format_table(table)


def add_forecast_table(doc: Document, stats: dict) -> None:
    summary = stats["submission_summary"]
    rows = [
        ("Горизонт прогноза", f"{summary['forecast_month']}-й месяц"),
        ("Количество строк в test.csv", fmt_int(summary["rows"])),
        ("Колонки", ", ".join(summary["columns"])),
        ("Минимальный прогноз РТО", fmt_money(summary["rto_min"])),
        ("Медианный прогноз РТО", fmt_money(summary["rto_median"])),
        ("Максимальный прогноз РТО", fmt_money(summary["rto_max"])),
        ("Суммарный прогноз сети", f"{summary['rto_sum'] / 1_000_000_000:,.2f}".replace(",", " ") + " млрд"),
        ("Основная модель", summary["primary_model"]),
        ("Резервная модель", summary["backup_model"]),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_cell_text(table.rows[0].cells[0], "Показатель", bold=True)
    set_cell_text(table.rows[0].cells[1], "Значение", bold=True)
    for left, right in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], left)
        set_cell_text(cells[1], right)
    format_table(table)


def add_recommendation_table(doc: Document) -> None:
    rows = [
        ("Планирование запасов", "Использовать прогноз РТО 11-го месяца как вход для закупок и поставок по магазинам", "Снизить риск дефицита в магазинах с растущим прогнозом и избыточных запасов в слабых точках"),
        ("Контроль отклонений", "Еженедельно сравнивать фактический РТО с прогнозом и выделять магазины с отклонением выше 8-10%", "Быстрее находить локальные проблемы: трафик, ассортимент, конкуренты, операционные сбои"),
        ("Фокус на лаги продаж", "Для магазинов с резким ростом/падением последних месяцев проводить отдельную проверку причин", "Главный фактор модели - РТО прошлого месяца, поэтому динамика последних периодов критична"),
        ("Региональные решения", "Сравнивать средний прогноз по регионам и корректировать региональные планы продаж", "Регионы с высоким средним прогнозом получают больший приоритет в маркетинговых активностях"),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ["Направление", "Рекомендация", "Ожидаемый эффект"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    format_table(table)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)


def build_report() -> None:
    stats = load_json(OUTPUTS_DIR / "report_stats.json")
    metrics = pd.read_csv(OUTPUTS_DIR / "metrics_saved_models.csv")

    doc = Document()
    configure_document(doc)

    # The user asked to keep the first sheet empty for the title page.
    # A tiny white marker prevents renderers from dropping the blank first page.
    blank_cover = doc.add_paragraph()
    blank_cover.paragraph_format.space_after = Pt(0)
    marker = blank_cover.add_run(".")
    marker.font.size = Pt(1)
    marker.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Лабораторная работа № 4")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Применение методов машинного обучения для прогнозирования РТО")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)

    add_paragraph(doc, "Первый лист оставлен пустым под титульный лист по требованию задания.")
    p = add_paragraph(doc, "Ссылка на код: ")
    add_hyperlink(p, stats["code_url"], stats["code_url"])
    p = add_paragraph(doc, "Основной скрипт без обучения моделей: ")
    add_hyperlink(p, stats["script_url"], stats["script_url"])

    add_heading(doc, "1. Постановка бизнес-задачи")
    add_paragraph(
        doc,
        "Бизнес-задача: для сети магазинов требуется спрогнозировать РТО на следующий, 11-й месяц по каждому магазину. "
        "Прогноз нужен для планирования запасов, логистики, региональных планов продаж и раннего выявления магазинов с ожидаемым отклонением оборота.",
    )
    add_paragraph(
        doc,
        "ML-задача: регрессия, поскольку целевая переменная является непрерывным числовым значением. "
        "Объект прогноза - магазин в конкретном прогнозном месяце, целевой показатель - РТО.",
    )
    add_bullets(
        doc,
        [
            f"Исходный период: месяцы {stats['month_min']}-{stats['month_max']}, всего {stats['month_count']} исторических месяцев.",
            f"Количество магазинов: {fmt_int(stats['store_count'])}; количество регионов: {fmt_int(stats['region_count'])}.",
            "Финальная сборка не обучает модели: используются сохраненные LinearRegression, RandomForestRegressor и TorchMLP.",
        ],
    )

    add_heading(doc, "2. Паспорт набора данных")
    add_passport_table(doc, stats)

    add_heading(doc, "3. Исследовательский анализ данных")
    add_paragraph(
        doc,
        f"В исходном наборе {fmt_int(stats['raw_rows'])} строк и {stats['raw_cols']} столбцов. "
        f"Медианное значение РТО составляет {fmt_money(stats['target_median'])}, среднее - {fmt_money(stats['target_mean'])}, "
        f"максимальное наблюдение - {fmt_money(stats['target_max'])}. Пропуски не обнаружены.",
    )
    add_paragraph(
        doc,
        "Распределение РТО имеет правый хвост: часть магазинов заметно крупнее медианного магазина. "
        "Такие наблюдения не удалялись, потому что они отражают реальный масштаб бизнеса и важны для прогноза.",
    )
    add_picture(doc, "eda_target_distribution.png", "Рисунок 1 - распределение исторического РТО")
    add_picture(doc, "eda_monthly_history.png", "Рисунок 2 - динамика суммарного и среднего РТО по месяцам")
    add_picture(doc, "eda_correlation_heatmap.png", "Рисунок 3 - корреляции ключевых признаков с РТО")
    add_picture(doc, "eda_category_profile.png", "Рисунок 4 - средний РТО по категориальному профилю магазина")
    add_paragraph(
        doc,
        "Сильнейшая связь с целевым показателем ожидаемо наблюдается у лагов РТО: РТО прошлого месяца, среднего РТО за предыдущие периоды и РТО два месяца назад. "
        "Это подтверждает, что задача носит временной характер, а последние месяцы продаж являются ключевым источником сигнала.",
    )

    add_heading(doc, "4. Предобработка данных")
    add_preprocessing_table(doc)
    add_paragraph(
        doc,
        f"После формирования лагов и удаления первых месяцев без достаточной истории осталось {fmt_int(stats['processed_rows'])} строк. "
        "Для месяца были добавлены циклические признаки sin/cos, чтобы модель учитывала сезонность без разрыва между декабрем и январем.",
    )

    add_heading(doc, "5. Построение моделей")
    add_paragraph(
        doc,
        "В финальной папке код не содержит вызовов обучения. Сравнение выполнено путем загрузки сохраненных моделей и расчета прогнозов на holdout-выборке. "
        "Это соответствует требованию использовать уже сохраненные и натренированные модели.",
    )
    add_models_table(doc)
    add_paragraph(doc, "Фрагмент финального пайплайна инференса:")
    add_code_block(
        doc,
        [
            "linear_model = joblib.load(SAVE_MODELS_DIR / 'linear_regression.pkl')",
            "rf_model = joblib.load(SAVE_MODELS_DIR / 'random_forest.pkl')",
            "checkpoint = torch.load(SAVE_MODELS_DIR / 'torch_mlp.pth', map_location='cpu')",
            "predictions = model(torch.tensor(X_scaled.to_numpy(), dtype=torch.float32))",
        ],
    )

    add_heading(doc, "6. Оценка качества моделей")
    add_metrics_table(doc, metrics)
    add_picture(doc, "model_metrics_comparison.png", "Рисунок 5 - сравнение качества сохраненных моделей")
    best = stats["best_model"]
    add_paragraph(
        doc,
        f"Лучшая модель по MAPE - {best['model']} со значением {fmt_pct(best['MAPE'])}. "
        f"MAE составляет {fmt_money(best['MAE'])}, RMSE - {fmt_money(best['RMSE'])}, R² - {fmt_metric(best['R2'], 4)}.",
    )

    add_heading(doc, "7. Анализ важности признаков")
    add_picture(doc, "feature_importance_random_forest.png", "Рисунок 6 - топ признаков по Random Forest")
    add_paragraph(
        doc,
        "Наиболее важный признак - РТО прошлого месяца. Это логично: оборот магазина устойчив во времени, а краткосрочный прогноз в первую очередь зависит от последних продаж. "
        "Также важны средние лаговые показатели магазина, сезонные признаки month_sin/month_cos, рост относительно прошлого месяца и параметры локального рынка.",
    )

    doc.add_page_break()
    add_heading(doc, "8. Итоговый прогноз")
    add_forecast_table(doc, stats)
    add_picture(doc, "forecast_distribution.png", "Рисунок 7 - распределение прогноза РТО на 11-й месяц")
    add_picture(doc, "network_total_history_forecast.png", "Рисунок 8 - история сети и прогноз на 11-й месяц")
    add_picture(doc, "top_regions_average_forecast.png", "Рисунок 9 - регионы с наибольшим средним прогнозом")
    add_picture(doc, "october_vs_forecast.png", "Рисунок 10 - сравнение фактического октября и прогноза ноября")

    add_heading(doc, "9. Бизнес-выводы и рекомендации")
    add_recommendation_table(doc)
    add_paragraph(
        doc,
        "Практический смысл модели - не только получить файл прогноза, но и превратить прогноз в управленческий контур: план поставок, мониторинг отклонений, региональные цели и приоритизация магазинов с резкой динамикой.",
    )

    add_heading(doc, "10. Заключение")
    add_bullets(
        doc,
        [
            "Сформулирована задача регрессии для прогноза РТО магазинов на 11-й месяц.",
            "Проведен EDA: данные полные, целевой признак асимметричен, сильнейший сигнал дают лаговые признаки РТО.",
            "Подготовлены признаки: лаги, скользящие средние, темп роста, one-hot регионов, ordinal-кодирование категорий и масштабирование числовых признаков.",
            "Сравнены сохраненные модели LinearRegression, RandomForestRegressor и TorchMLP без повторного обучения.",
            "По MAPE лучшей стала TorchMLP; она использована как основная модель для итогового файла test.csv.",
        ],
    )
    add_paragraph(
        doc,
        "Ограничения: модель использует только 10 месяцев истории и не содержит внешних факторов вроде промоакций, праздников, цен конкурентов и погодных условий. "
        "Перспективы улучшения: добавить внешние календарные признаки, промо-факторы, контроль качества по регионам и мониторинг drift после получения факта 11-го месяца.",
    )

    add_heading(doc, "11. Ответы на контрольные вопросы")
    add_numbered(
        doc,
        [
            "Регрессия прогнозирует число, например объем продаж или стоимость заказа; классификация выбирает класс, например факт оттока клиента или принадлежность к сегменту. В данной работе используется регрессия, потому что РТО является непрерывной величиной.",
            "Категориальные признаки кодируются, потому что большинство ML-алгоритмов работает с числами. One-Hot Encoding подходит для категорий без порядка, например регионов; Ordinal Encoding подходит для упорядоченных категорий, например размера торговой площади.",
            "При 10% ушедших клиентов Accuracy может быть обманчивой: модель, предсказывающая всем отсутствие оттока, даст 90% Accuracy. Нужны Precision, Recall, F1-score и ROC-AUC; для удержания клиентов часто особенно важны Recall и F1.",
            "Feature Importance в Random Forest показывает вклад признака в качество разбиений деревьев. В задаче оттока телеком-клиентов важными могли бы быть длительность контракта, число обращений в поддержку и размер ежемесячного платежа; бизнес использует это для таргетированных удерживающих действий.",
            "Рост R² с 0,85 до 0,90 означает, что признак расстояния до метро объясняет дополнительную долю вариации цены и практически полезен для оценки объектов. В другой бизнес-задаче R² полезен, например, при прогнозировании выручки магазина, где важна объясненная доля разброса продаж.",
        ],
    )

    doc.save(REPORT_PATH)


if __name__ == "__main__":
    build_report()
