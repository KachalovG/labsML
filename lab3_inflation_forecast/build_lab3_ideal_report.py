from __future__ import annotations

import json
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SAVE_DIR = ROOT / "models" / "save_models"
ASSETS_DIR = ROOT / "report_assets"
REPORTS_DIR = ROOT / "reports"
FINAL_REPORT = REPORTS_DIR / "ЛР3_Качалов_идеальный_отчет.docx"
DESKTOP_COPY = Path(r"C:\Users\gvkac\Desktop\учеба\ЛР\ЛР №3 МАП готовый отчет Качалов.docx")
REPOSITORY_URL = "https://github.com/KachalovG/labsML/tree/main/lab3_inflation_forecast"

MONTHS = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_cell_padding(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    return doc


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_para(doc: Document, text: str = "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.alignment = align
    set_paragraph_spacing(p)
    if bold_prefix and text.startswith(bold_prefix):
        prefix = p.add_run(bold_prefix)
        set_run_font(prefix, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc: Document, items):
    for item in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        set_paragraph_spacing(p, after=3)
        run = p.add_run(f"• {item}")
        set_run_font(run)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=8)
    run = p.add_run(text)
    set_run_font(run, size=10, italic=True)
    return p


def add_table_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=6, after=4)
    keep_with_next(p)
    run = p.add_run(text)
    set_run_font(run, size=10, bold=True)
    return p


def add_table(doc: Document, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.05)
        run = p.add_run(str(header))
        set_run_font(run, size=font_size, bold=True)
        set_cell_shading(cell, "D9EAF7")
        set_cell_padding(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(cell, widths[idx])

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            text = str(value)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, after=0, line=1.05)
            run = p.add_run(text)
            set_run_font(run, size=font_size)
            set_cell_padding(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[idx])

    doc.add_paragraph()
    return table


def add_picture(doc: Document, image_name: str, caption: str, width_cm=15.5):
    path = ASSETS_DIR / image_name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def fmt(value, digits=4):
    if pd.isna(value):
        return ""
    return f"{float(value):.{digits}f}"


def date_label(value):
    ts = pd.to_datetime(value)
    return f"{MONTHS[ts.month].lower()} {ts.year}"


def load_sources():
    raw = pd.read_csv(ROOT / "data" / "inflation_long_format.csv", parse_dates=["date"])
    metrics = pd.read_csv(SAVE_DIR / "model_metrics.csv")
    seasonality = pd.read_csv(SAVE_DIR / "seasonality.csv")
    forecast = pd.read_csv(SAVE_DIR / "future_forecast.csv", parse_dates=["date"])
    predictions = pd.read_csv(SAVE_DIR / "predictions.csv", parse_dates=["date"])
    with open(SAVE_DIR / "features.json", encoding="utf-8") as fh:
        features = json.load(fh)
    return raw, metrics, seasonality, forecast, predictions, features


def add_title_page(doc: Document):
    centered_lines = [
        "Министерство науки и высшего образования Российской Федерации",
        "Федеральное государственное автономное образовательное учреждение высшего образования",
        "«Московский государственный технический университет имени Н.Э. Баумана»",
        "(МГТУ им. Н.Э. Баумана)",
        "",
        "ФАКУЛЬТЕТ «ИНЖЕНЕРНЫЙ БИЗНЕС И МЕНЕДЖМЕНТ»",
        "КАФЕДРА «ПРОМЫШЛЕННАЯ ЛОГИСТИКА» (ИБМ-3)",
    ]
    for line in centered_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=3)
        run = p.add_run(line)
        set_run_font(run, size=12, bold=line.startswith(("ФАКУЛЬТЕТ", "КАФЕДРА")))

    for _ in range(3):
        doc.add_paragraph()

    for text, size, bold in [
        ("Отчет по лабораторной работе № 3", 16, True),
        ("по дисциплине «Методы аналитики и прогнозирования»", 14, False),
        ("«Построение модели прогноза инфляции с сезонной компонентой»", 14, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=8)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)

    for _ in range(3):
        doc.add_paragraph()

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, row in enumerate([["Студент", "ИБМ 3-55Б    Г.В. Качалов"], ["Руководитель", "Е.Д. Мелихова"]]):
        for col_idx, text in enumerate(row):
            cell = table.rows[idx].cells[col_idx]
            set_cell_width(cell, 4 if col_idx == 0 else 11)
            set_cell_padding(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, after=0)
            run = p.add_run(text)
            set_run_font(run, size=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=24, after=2)
    run = p.add_run("Москва, 2026")
    set_run_font(run)
    doc.add_page_break()


def add_header_footer(doc: Document):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(header, after=0)
    run = header.add_run("Лабораторная работа №3. Прогноз инфляции с сезонной компонентой")
    set_run_font(run, size=9, color=(89, 89, 89))

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        set_run_font(run, size=9, color=(89, 89, 89))


def build_report():
    raw, metrics, seasonality, forecast, predictions, features = load_sources()
    REPORTS_DIR.mkdir(exist_ok=True)

    doc = setup_document()
    add_header_footer(doc)
    add_title_page(doc)

    raw_min = raw["date"].min().strftime("%Y-%m-%d")
    raw_max = raw["date"].max().strftime("%Y-%m-%d")
    total_raw = raw[raw["inflation_type"] == "total"].copy()
    max_total = total_raw.loc[total_raw["value"].idxmax()]
    min_total = total_raw.loc[total_raw["value"].idxmin()]

    add_heading(doc, "1. Цель работы")
    add_para(
        doc,
        "Цель работы — сформировать практические навыки анализа временных рядов, "
        "выделения сезонной компоненты, построения прогнозных моделей и оценки "
        "точности прогноза для поддержки управленческих решений.",
    )
    add_para(
        doc,
        "По согласованию с преподавателем вместо Excel-реализации были построены "
        "модели в Python. Требование задания о сезонной компоненте сохранено: "
        "в работе рассчитаны коэффициенты сезонности, использованы месячные признаки, "
        "лаговые признаки, скользящие средние и сезонные baseline-модели.",
        bold_prefix="По согласованию с преподавателем",
    )

    add_table_caption(doc, "Таблица 1 — Соответствие отчета требованиям лабораторной работы")
    add_table(
        doc,
        ["Требование задания", "Что выполнено в отчете"],
        [
            ["Выбор временного ряда не менее чем за 3 года", f"Использован ряд месячной инфляции РФ за период {raw_min} — {raw_max}."],
            ["Первичная визуализация и оценка тренда/сезонности", "Построен график полного ряда и современного периода, выделены выбросы и сезонные месячные эффекты."],
            ["Два способа расчета сезонности", "Рассчитаны коэффициенты через средние по годам и через сводный среднемесячный профиль."],
            ["Построение прогнозной модели", "Построены Linear Regression, Ridge, HuberRegressor, Gradient Boosting и четыре baseline-модели."],
            ["Оценка точности", "Рассчитаны MAE, MSE, RMSE, R2 и MAPE по индексам для validation, test, holdout и recent actual."],
            ["Итоговый прогноз", "Сформирован прогноз на 12 месяцев вперед: апрель 2026 — март 2027."],
            ["Приложение с расчетами", "В проекте сохранены исходные данные, код, метрики, прогнозы, графики и обученная модель; ссылка на GitHub приведена в приложении."],
        ],
        widths=[6, 10],
        font_size=9,
    )

    h = add_heading(doc, "2. Описание исходных данных")
    h.paragraph_format.page_break_before = True
    add_para(
        doc,
        "В качестве предметной области выбран макроэкономический временной ряд: "
        "индексы потребительских цен по Российской Федерации. Для моделирования "
        "использовались четыре категории: продовольственные товары, непродовольственные "
        "товары, услуги и общий индекс.",
    )
    add_table_caption(doc, "Таблица 2 — Паспорт исходных данных")
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Наименование показателя", "Индексы потребительских цен; месячная инфляция"],
            ["Источник", "Локальный файл data/inflation_long_format.csv; первичный источник данных — Росстат"],
            ["Период данных", f"{raw_min} — {raw_max}"],
            ["Количество наблюдений", f"{len(raw)} строк в исходном файле; {features['train_rows'] + features['val_rows'] + features['test_rows'] + features['holdout_rows'] + features['recent_rows']} строк после формирования лагов"],
            ["Частота наблюдений", "Ежемесячно"],
            ["Категории", "food, non_food, services, total"],
            ["Целевая переменная", "target = idx - 100, где idx — индекс цен к предыдущему месяцу"],
            ["Пропуски", f"{int(raw.isna().sum().sum())} пропусков"],
            ["Репозиторий с кодом", REPOSITORY_URL],
        ],
        widths=[5, 11],
        font_size=9.5,
    )

    add_heading(doc, "3. Первичный анализ временного ряда")
    add_para(
        doc,
        f"Полный ряд содержит экстремальные значения начала 1990-х годов. Максимальный "
        f"общий индекс зафиксирован в {date_label(max_total['date'])}: {fmt(max_total['value'], 2)}. "
        f"Минимальное значение общего индекса в доступных данных — {fmt(min_total['value'], 2)} "
        f"({date_label(min_total['date'])}). Такие выбросы важны: они ухудшают качество простых "
        f"линейных моделей и требуют либо робастных методов, либо нелинейных алгоритмов.",
    )
    add_picture(
        doc,
        "time_series.png",
        "Рисунок 1 — Динамика месячной инфляции по категориям за весь период и за современный период",
        width_cm=15.7,
    )
    add_para(
        doc,
        "На современном участке после 2016 года колебания находятся в более узком диапазоне, "
        "но остаются отдельные всплески. Поэтому модель должна одновременно учитывать сезонность, "
        "лаговую инерцию и устойчивость к аномальным периодам.",
    )
    add_picture(
        doc,
        "target_distribution.png",
        "Рисунок 2 — Распределение целевой переменной target = idx - 100",
        width_cm=13.8,
    )

    add_heading(doc, "4. Подготовка данных")
    add_para(doc, "Подготовка данных выполнена программно в модуле data/data_processing.py. Основные шаги:")
    add_bullets(
        doc,
        [
            "дата приведена к месячному формату, каждая строка соответствует одному месяцу и одной категории инфляции;",
            "месяц закодирован one-hot признаками m_01 ... m_12;",
            "категория инфляции закодирована one-hot признаками type_food, type_non_food, type_services, type_total;",
            "рассчитаны лаги lag_1, lag_2, lag_3, lag_6, lag_12;",
            "рассчитаны скользящие средние roll_3, roll_6, roll_12;",
            "числовые признаки стандартизированы с помощью StandardScaler;",
            "целевая переменная определена как отклонение индекса от 100: target = idx - 100.",
        ],
    )
    add_table_caption(doc, "Таблица 3 — Разбиение данных на выборки")
    add_table(
        doc,
        ["Выборка", "Период", "Назначение", "Строк"],
        [
            ["Train", "1992-2018", "Обучение моделей", features["train_rows"]],
            ["Validation", "2019-2022", "Выбор лучшей trainable-модели", features["val_rows"]],
            ["Test", "2023-2024", "Финальная проверка качества", features["test_rows"]],
            ["Holdout", "2025", "Проверка устойчивости на новом годе", features["holdout_rows"]],
            ["Recent actual", "2026", "Последние доступные фактические значения", features["recent_rows"]],
        ],
        widths=[3, 3, 7, 2.2],
        font_size=9.5,
    )

    h = add_heading(doc, "5. Расчет коэффициентов сезонности")
    h.paragraph_format.page_break_before = True
    add_para(
        doc,
        "Сезонность рассчитана двумя способами. Первый способ нормирует значение месяца "
        "на среднее значение соответствующего года. Второй способ сопоставляет среднемесячное "
        "значение с общим средним по всем месяцам. Для месячной сезонности сумма коэффициентов "
        "должна быть близка к 12; в расчетах для общего индекса она составляет 11,9983 и 11,9979.",
    )
    add_para(doc, "Использованные формулы:")
    add_bullets(
        doc,
        [
            "КС(месяц, год) = Y(месяц, год) / среднее значение за соответствующий год;",
            "КС(месяц) = среднее значение месяца / общее среднее значение ряда;",
            "значение КС > 1 означает месяц выше среднего уровня, КС < 1 — месяц ниже среднего уровня.",
        ],
    )
    total_season = seasonality[seasonality["inflation_type"] == "total"].copy()
    season_rows = []
    for row in total_season.itertuples(index=False):
        season_rows.append(
            [
                MONTHS[int(row.month_num)],
                fmt(row.seasonality_by_year_method, 4),
                fmt(row.seasonality_pivot_method, 4),
                fmt(row.month_mean_idx, 2),
            ]
        )
    add_table_caption(doc, "Таблица 4 — Коэффициенты сезонности для общей инфляции")
    add_table(
        doc,
        ["Месяц", "Метод средних по годам", "Метод сводного профиля", "Средний idx"],
        season_rows,
        widths=[3.3, 4.2, 4.2, 3],
        font_size=9,
    )
    add_picture(doc, "seasonality.png", "Рисунок 3 — Сезонный профиль общей инфляции", width_cm=14.8)
    add_para(
        doc,
        "Наиболее выраженный сезонный эффект наблюдается в январе: после пересмотра цен и тарифов "
        "индекс в среднем выше общего уровня. Самые спокойные месяцы по общей инфляции — летние, "
        "особенно август. Для инфляции сезонная амплитуда слабее, чем для товарных продаж, но она "
        "остается полезной для прогноза как объясняющий признак.",
    )

    h = add_heading(doc, "6. Построение прогнозных моделей")
    h.paragraph_format.page_break_before = True
    add_para(
        doc,
        "Так как была разрешена реализация моделей в Python, классическая Excel-схема "
        "«Лист прогноза / ручной Хольт-Винтерс» заменена на набор моделей машинного обучения "
        "и сезонных базовых моделей. Это расширяет сравнение: ML-модель должна быть лучше "
        "простого прогноза по прошлому месяцу, прошлому году или скользящему среднему.",
    )
    add_table_caption(doc, "Таблица 5 — Использованные прогнозные модели")
    add_table(
        doc,
        ["Модель", "Тип", "Роль в работе", "Учет сезонности"],
        [
            ["naive_lag_1", "Baseline", "Прогноз равен значению прошлого месяца", "Лаговая инерция"],
            ["naive_lag_12", "Seasonal baseline", "Прогноз равен значению того же месяца прошлого года", "Явная годовая сезонность"],
            ["rolling_3 / rolling_12", "Baseline", "Прогноз по скользящему среднему", "Сглаживание кратко- и среднесрочной динамики"],
            ["Linear Regression", "ML", "Интерпретируемая линейная модель", "Месячные one-hot признаки и лаги"],
            ["Ridge", "ML", "Линейная модель с L2-регуляризацией", "Месячные one-hot признаки и лаги"],
            ["HuberRegressor", "Robust ML", "Устойчивая линейная модель для выбросов", "Месячные признаки, лаги, робастная функция потерь"],
            ["GradientBoostingRegressor", "Nonlinear ML", "Итоговая модель по validation/test", "Нелинейное взаимодействие месяцев, лагов и категорий"],
        ],
        widths=[3.4, 2.6, 5.2, 4.4],
        font_size=8.5,
    )
    add_para(
        doc,
        "Для обучаемых моделей подбор гиперпараметров выполнен через GridSearchCV с "
        "TimeSeriesSplit(n_splits=3), что корректнее случайного перемешивания для временного ряда. "
        "Итоговая модель Gradient Boosting использует loss='huber', n_estimators=200, "
        "learning_rate=0.03, max_depth=2, min_samples_leaf=4, random_state=42.",
    )
    add_picture(
        doc,
        "feature_importance.png",
        "Рисунок 4 — Важность признаков итоговой модели Gradient Boosting",
        width_cm=14.2,
    )

    add_heading(doc, "7. Оценка точности")
    add_para(
        doc,
        "Качество оценивалось на отложенных временных периодах. Основные метрики: "
        "MAE — средняя абсолютная ошибка в процентных пунктах target; MSE и RMSE сильнее "
        "штрафуют крупные ошибки; R2 показывает объясненную долю вариации; MAPE рассчитан "
        "по индексам idx, чтобы избежать нестабильности процентной ошибки при target около нуля.",
    )
    test_metrics = metrics[metrics["dataset"] == "test_2023_2024"].sort_values("mae")
    metric_rows = [
        [r.model, fmt(r.mae), fmt(r.mse), fmt(r.rmse), fmt(r.r2), fmt(r.mape_idx)]
        for r in test_metrics.itertuples(index=False)
    ]
    add_table_caption(doc, "Таблица 6 — Сравнение точности моделей на test 2023-2024")
    add_table(
        doc,
        ["Модель", "MAE", "MSE", "RMSE", "R2", "MAPE, %"],
        metric_rows,
        widths=[4, 2.1, 2.1, 2.1, 2.1, 2.1],
        font_size=8.7,
    )
    add_picture(doc, "test_mae.png", "Рисунок 5 — Сравнение моделей на тестовом периоде 2023-2024 по MAE", width_cm=14.4)
    add_picture(doc, "actual_vs_pred.png", "Рисунок 6 — Сопоставление фактической и прогнозной общей инфляции", width_cm=14.8)

    holdout_metrics = metrics[metrics["dataset"] == "holdout_2025"].sort_values("mae").head(5)
    holdout_rows = [
        [r.model, fmt(r.mae), fmt(r.rmse), fmt(r.r2), fmt(r.mape_idx)]
        for r in holdout_metrics.itertuples(index=False)
    ]
    add_table_caption(doc, "Таблица 7 — Проверка устойчивости на holdout 2025")
    add_table(
        doc,
        ["Модель", "MAE", "RMSE", "R2", "MAPE, %"],
        holdout_rows,
        widths=[4.2, 2.2, 2.2, 2.2, 2.2],
        font_size=9,
    )
    add_para(
        doc,
        "На validation и test лучшей обучаемой моделью стала Gradient Boosting. На holdout 2025 "
        "наименьшую ошибку показал HuberRegressor, что дополнительно подтверждает чувствительность "
        "экономических рядов к выбросам. Для итогового прогноза выбрана Gradient Boosting как модель, "
        "победившая на этапе выбора и финального тестирования, а результат holdout зафиксирован как "
        "ограничение интерпретации.",
    )

    h = add_heading(doc, "8. Итоговый прогноз")
    h.paragraph_format.page_break_before = True
    total_forecast = forecast[forecast["inflation_type"] == "total"].sort_values("date")
    forecast_rows = [
        [date_label(r.date), fmt(r.pred_target), fmt(r.pred_idx)]
        for r in total_forecast.itertuples(index=False)
    ]
    add_table_caption(doc, "Таблица 8 — Итоговый прогноз общей инфляции на 12 месяцев")
    add_table(
        doc,
        ["Период", "Прогноз target, п.п.", "Прогноз idx"],
        forecast_rows,
        widths=[5, 4.6, 3.6],
        font_size=9.2,
    )
    add_picture(doc, "future_forecast.png", "Рисунок 7 — Итоговый прогноз общей инфляции на 12 месяцев", width_cm=14.8)
    peak = total_forecast.loc[total_forecast["pred_target"].idxmax()]
    calm = total_forecast.loc[total_forecast["pred_target"].idxmin()]
    add_para(
        doc,
        f"По прогнозу максимальная месячная инфляция ожидается в {date_label(peak['date'])}: "
        f"{fmt(peak['pred_target'], 2)} п.п. Минимальное прогнозное значение — "
        f"{fmt(calm['pred_target'], 2)} п.п. в {date_label(calm['date'])}. Для управленческого "
        f"планирования это означает, что индексацию цен, пересмотр бюджетов и закупочные решения "
        f"целесообразно готовить до месяцев с прогнозируемым ускорением инфляции.",
    )

    add_heading(doc, "9. Бизнес-интерпретация")
    add_bullets(
        doc,
        [
            "Финансовое планирование: прогноз можно использовать для пересмотра бюджета расходов и планирования индексации договоров.",
            "Закупки и запасы: перед ожидаемыми пиковыми месяцами можно заранее фиксировать цены и корректировать график закупок.",
            "Ценообразование: прогноз общей инфляции помогает обосновать календарь пересмотра цен и тарифов.",
            "Управление рисками: расхождение качества на test и holdout показывает, что модель нужно регулярно переобучать после появления новых фактических месяцев.",
        ],
    )

    add_heading(doc, "10. Выводы")
    add_bullets(
        doc,
        [
            "Временной ряд инфляции подготовлен к моделированию: сформированы лаги, скользящие средние, сезонные и категориальные признаки.",
            "Коэффициенты сезонности рассчитаны двумя методами; для общего индекса наиболее заметен январский эффект.",
            "Построены и сравнены несколько моделей, включая простые сезонные baseline и обучаемые ML-модели.",
            "По test 2023-2024 лучшая точность у Gradient Boosting: MAE = 0,3699, MAPE = 0,3666%.",
            "На holdout 2025 наиболее устойчивым оказался HuberRegressor; это зафиксировано как ограничение итогового прогноза.",
            "Итоговый прогноз построен на 12 месяцев вперед и может использоваться для бюджетирования, закупок, ценообразования и оценки инфляционных рисков.",
        ],
    )

    add_heading(doc, "11. Ответы на контрольные вопросы")
    qa = [
        (
            "1. Компоненты временного ряда.",
            "Временной ряд включает тренд, сезонную и случайную компоненты. Тренд отражает "
            "долгосрочное направление изменения, сезонность — повторяющиеся колебания внутри "
            "года, случайная компонента — нерегулярные отклонения. Наличие сезонности требует "
            "модели, которая учитывает месячный или квартальный цикл."
        ),
        (
            "2. Коэффициент сезонности.",
            "Коэффициент сезонности показывает, насколько конкретный месяц отличается от "
            "среднего уровня. Если коэффициент декабря равен 1,25, это означает, что показатель "
            "в декабре в среднем на 25% выше обычного уровня, поэтому бизнесу нужно заранее "
            "готовить запасы, персонал или бюджет."
        ),
        (
            "3. Аддитивная и мультипликативная модели.",
            "В аддитивной модели сезонная добавка примерно постоянна в абсолютных единицах, "
            "а в мультипликативной сезонность пропорциональна уровню ряда. Если амплитуда "
            "сезонных колебаний растет вместе с уровнем показателя, лучше подходит "
            "мультипликативная логика."
        ),
        (
            "4. MAPE.",
            "MAPE — средняя абсолютная процентная ошибка. Значение ниже 10% обычно считается "
            "высокой точностью, 10-20% — хорошей, 20-50% — удовлетворительной. При очень малых "
            "фактических значениях MAPE может вводить в заблуждение, поэтому в работе ошибка "
            "рассчитывалась по индексам idx."
        ),
        (
            "5. Использование прогноза.",
            "Прогноз можно использовать в финансах, закупках и ценообразовании. Финансовый блок "
            "планирует бюджет и индексацию, закупки выбирают момент фиксации цен и объемы, "
            "а коммерческий блок определяет календарь пересмотра цен."
        ),
    ]
    for question, answer in qa:
        add_para(doc, f"{question} {answer}", bold_prefix=question)

    add_heading(doc, "12. Приложение")
    add_para(doc, "Проект содержит код подготовки данных, обучения моделей, построения графиков и генерации отчета.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p)
    run = p.add_run("Репозиторий с кодом: ")
    set_run_font(run)
    add_hyperlink(p, REPOSITORY_URL, REPOSITORY_URL)
    add_bullets(
        doc,
        [
            "data/inflation_long_format.csv — исходный датасет;",
            "data/data_processing.py — подготовка признаков;",
            "train.py — обучение, расчет метрик, сохранение прогнозов;",
            "models/ — определения моделей;",
            "models/save_models/ — метрики, прогнозы, коэффициенты сезонности, обученная модель;",
            "report_assets/ — графики для отчета;",
            "build_lab3_ideal_report.py — генерация финального Word-отчета.",
        ],
    )

    doc.save(FINAL_REPORT)
    DESKTOP_COPY.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(FINAL_REPORT, DESKTOP_COPY)
    return FINAL_REPORT, DESKTOP_COPY


if __name__ == "__main__":
    report, copy = build_report()
    print(report)
    print(copy)
