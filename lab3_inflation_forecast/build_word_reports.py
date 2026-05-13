import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from data.data_processing import prepare_data


ROOT = Path(".")
SAVE_DIR = ROOT / "models" / "save_models"
ASSETS_DIR = ROOT / "report_assets"
REPORTS_DIR = ROOT / "reports"

LR3_PATH = REPORTS_DIR / "ЛР3_Качалов_прогноз_инфляции.docx"
LR4_PATH = REPORTS_DIR / "ЛР4_Качалов_ML_инфляция.docx"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_cm):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(header_cells[i], header, bold=True)
        set_cell_shading(header_cells[i], "D9EAF7")
        if widths:
            set_cell_width(header_cells[i], widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if widths:
                set_cell_width(cells[i], widths[i])

    doc.add_paragraph()
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_table_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_picture(doc, image_name, caption, width_cm=15.5):
    path = ASSETS_DIR / image_name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(31, 78, 121)
    return p


def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = p.add_run(bold_prefix)
        prefix_run.bold = True
        prefix_run.font.name = "Times New Roman"
        rest = text[len(bold_prefix):]
        run = p.add_run(rest)
    else:
        run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(31, 78, 121)

    return doc


def add_title_page(doc, lab_number, topic):
    centered = [
        "Министерство науки и высшего образования Российской Федерации",
        "Федеральное государственное автономное образовательное учреждение высшего образования",
        "«Московский государственный технический университет имени Н.Э. Баумана»",
        "(МГТУ им. Н.Э. Баумана)",
        "",
        "ФАКУЛЬТЕТ «ИНЖЕНЕРНЫЙ БИЗНЕС И МЕНЕДЖМЕНТ»",
        "КАФЕДРА «ПРОМЫШЛЕННАЯ ЛОГИСТИКА» (ИБМ-3)",
    ]
    for text in centered:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        if "ФАКУЛЬТЕТ" in text or "КАФЕДРА" in text:
            run.bold = True

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Отчет по лабораторной работе № {lab_number}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("по дисциплине «Методы аналитики и прогнозирования»")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(topic)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    for _ in range(4):
        doc.add_paragraph()

    add_table(
        doc,
        ["", ""],
        [
            ["Студент", "ИБМ 3-55Б    Г.В. Качалов"],
            ["Руководитель", "Е.Д. Мелихова"]
        ],
        widths=[4, 11]
    )

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    doc.add_page_break()


def fmt(value, digits=4):
    if pd.isna(value):
        return ""
    if isinstance(value, (int, float)):
        return f"{value:.{digits}f}"
    return str(value)


def load_data():
    raw = pd.read_csv(ROOT / "data" / "inflation_long_format.csv")
    df = prepare_data(ROOT / "data" / "inflation_long_format.csv")
    type_cols = [col for col in df.columns if col.startswith("type_")]
    df["inflation_type"] = (
        df[type_cols]
        .idxmax(axis=1)
        .str.replace("type_", "", regex=False)
    )
    metrics = pd.read_csv(SAVE_DIR / "model_metrics.csv")
    predictions = pd.read_csv(SAVE_DIR / "predictions.csv")
    seasonality = pd.read_csv(SAVE_DIR / "seasonality.csv")
    future_forecast = pd.read_csv(SAVE_DIR / "future_forecast.csv")
    features = json.load(open(SAVE_DIR / "features.json", encoding="utf-8"))
    return raw, df, metrics, predictions, seasonality, future_forecast, features


def build_lr3():
    raw, df, metrics, predictions, seasonality, future, features = load_data()

    doc = setup_doc()
    add_title_page(
        doc,
        "3",
        "«Построение модели прогноза инфляции с сезонной компонентой»"
    )

    add_heading(doc, "1. Цель работы")
    add_paragraph(
        doc,
        "Цель работы — сформировать практические навыки анализа временного ряда, "
        "выделения сезонной компоненты, построения прогнозных моделей и оценки точности "
        "прогноза. В рамках данной работы вместо продаж рассматривается месячная инфляция "
        "в Российской Федерации; такая адаптация соответствует разрешенной Python-реализации."
    )

    add_heading(doc, "2. Описание исходных данных")
    add_table_caption(doc, "Таблица 1 — Паспорт исходных данных")
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Наименование показателя", "Индексы потребительских цен / месячная инфляция"],
            ["Источник", "Файл data/inflation_long_format.csv; первичный источник — Росстат"],
            ["Период данных", f"{raw['date'].min()} — {raw['date'].max()}"],
            ["Количество наблюдений", f"{len(raw)} строк в исходном файле; {len(df)} строк после формирования лагов"],
            ["Частота", "Ежемесячно"],
            ["Категории", ", ".join(sorted(raw["inflation_type"].unique()))],
            ["Целевая переменная", "target = idx - 100"],
            ["Пропуски", f"{int(raw.isna().sum().sum())} пропусков"]
        ],
        widths=[5, 10]
    )

    add_heading(doc, "3. Первичный анализ временного ряда")
    add_paragraph(
        doc,
        "На полном графике хорошо видны экстремальные значения начала 1990-х годов. "
        "Именно они объясняют, почему для современной инфляции обычные линейные модели "
        "работают хуже: модель частично подстраивается под исторические кризисные выбросы."
    )
    add_picture(
        doc,
        "time_series.png",
        "Рисунок 1 — Динамика месячной инфляции по категориям за весь период и за современный период"
    )
    add_paragraph(
        doc,
        "Отдельный график recent period нужен для анализа последних лет: на полном графике "
        "современные колебания почти незаметны из-за масштаба 1990-х. После 2016 года видно, "
        "что значения находятся в узком диапазоне, но присутствуют отдельные всплески, например в 2022 году."
    )

    add_heading(doc, "4. Подготовка данных")
    add_bullets(
        doc,
        [
            "Дата приведена к месячному формату; каждая строка соответствует одному месяцу и одной категории инфляции.",
            "Месяц закодирован one-hot признаками m_01 ... m_12.",
            "Категория инфляции закодирована one-hot признаками type_food, type_non_food, type_services, type_total.",
            "Рассчитаны лаги lag_1, lag_2, lag_3, lag_6, lag_12 и скользящие средние roll_3, roll_6, roll_12.",
            "Числовые признаки стандартизированы с помощью StandardScaler."
        ]
    )
    add_table_caption(doc, "Таблица 2 — Разбиение данных на выборки")
    add_table(
        doc,
        ["Выборка", "Период", "Назначение", "Строк"],
        [
            ["Train", "1992–2018", "Обучение моделей", features["train_rows"]],
            ["Validation", "2019–2022", "Выбор лучшей модели", features["val_rows"]],
            ["Test", "2023–2024", "Финальная проверка", features["test_rows"]],
            ["Holdout", "2025", "Проверка устойчивости", features["holdout_rows"]],
            ["Recent actual", "2026", "Последние доступные фактические значения", features["recent_rows"]]
        ],
        widths=[3, 3, 7, 2]
    )

    add_heading(doc, "5. Расчет коэффициентов сезонности")
    add_paragraph(
        doc,
        "Сезонность рассчитана двумя способами: через отношение месяца к среднему значению "
        "соответствующего года и через отношение среднемесячного значения к общему среднему."
    )
    seasonality_total = seasonality[seasonality["inflation_type"] == "total"].copy()
    add_table_caption(doc, "Таблица 3 — Коэффициенты сезонности для общей инфляции")
    add_table(
        doc,
        ["Месяц", "Метод средних по годам", "Метод сводной таблицы"],
        [
            [
                int(row["month_num"]),
                fmt(row["seasonality_by_year_method"], 4),
                fmt(row["seasonality_pivot_method"], 4)
            ]
            for _, row in seasonality_total.iterrows()
        ],
        widths=[3, 6, 6]
    )
    add_picture(
        doc,
        "seasonality.png",
        "Рисунок 2 — Сезонный профиль общей инфляции"
    )
    add_paragraph(
        doc,
        "Коэффициент выше 1 означает, что в данном месяце индекс инфляции в среднем выше "
        "общего уровня; коэффициент ниже 1 означает более спокойный месяц. Для инфляции "
        "сезонность слабее, чем для товарных продаж, но месячные эффекты полезны как признаки модели."
    )

    add_heading(doc, "6. Построение прогнозных моделей")
    add_paragraph(
        doc,
        "Вместо ручной Excel-реализации Хольта–Винтерса прогноз построен в Python. "
        "Сезонная логика сохранена через месячные признаки, лаги, скользящие средние и сезонные baseline-модели."
    )
    add_table_caption(doc, "Таблица 4 — Использованные прогнозные модели")
    add_table(
        doc,
        ["Модель", "Роль в работе"],
        [
            ["naive_lag_1", "Прогноз равен значению прошлого месяца"],
            ["naive_lag_12", "Прогноз равен значению того же месяца прошлого года"],
            ["rolling_3 / rolling_12", "Прогноз по скользящему среднему"],
            ["HuberRegressor", "Робастная модель, устойчивая к выбросам"],
            ["GradientBoostingRegressor", "Нелинейная модель для табличных временных признаков"]
        ],
        widths=[5, 10]
    )

    add_heading(doc, "7. Оценка точности")
    test_metrics = metrics[metrics["dataset"] == "test_2023_2024"].sort_values("mae")
    add_table_caption(doc, "Таблица 5 — Сравнение точности моделей на test 2023–2024")
    add_table(
        doc,
        ["Модель", "MAE", "MSE", "RMSE", "R2", "MAPE, %"],
        [
            [
                row["model"],
                fmt(row["mae"]),
                fmt(row["mse"]),
                fmt(row["rmse"]),
                fmt(row["r2"]),
                fmt(row["mape_idx"])
            ]
            for _, row in test_metrics.iterrows()
        ],
        widths=[4.2, 2.1, 2.1, 2.1, 2.1, 2.4]
    )
    add_picture(
        doc,
        "test_mae.png",
        "Рисунок 3 — Сравнение моделей на тестовом периоде 2023–2024 по MAE"
    )
    add_picture(
        doc,
        "actual_vs_pred.png",
        "Рисунок 4 — Сопоставление фактической и прогнозной общей инфляции"
    )
    add_paragraph(
        doc,
        "По validation и test лучшей моделью является Gradient Boosting. При этом на holdout 2025 "
        "самый устойчивый результат показал HuberRegressor, что подтверждает полезность робастных моделей "
        "для экономических рядов с выбросами."
    )

    add_heading(doc, "8. Итоговый прогноз")
    future_total = future[future["inflation_type"] == "total"].copy()
    add_table_caption(doc, "Таблица 6 — Итоговый прогноз общей инфляции на 12 месяцев")
    add_table(
        doc,
        ["Период", "Прогноз target", "Прогноз idx"],
        [
            [
                str(row["date"])[:10],
                fmt(row["pred_target"], 4),
                fmt(row["pred_idx"], 4)
            ]
            for _, row in future_total.iterrows()
        ],
        widths=[4, 5, 5]
    )
    add_picture(
        doc,
        "future_forecast.png",
        "Рисунок 5 — Итоговый прогноз общей инфляции на 12 месяцев"
    )

    add_heading(doc, "9. Выводы")
    add_bullets(
        doc,
        [
            "Были получены навыки подготовки месячного временного ряда, расчета лагов и сезонных коэффициентов.",
            "На данных инфляции сильные выбросы начала 1990-х существенно влияют на качество линейных моделей.",
            "Сезонные baseline-модели являются обязательной точкой сравнения: ML-модель должна превосходить простые лаги.",
            "Для итогового прогноза выбрана модель Gradient Boosting, а HuberRegressor отмечен как более устойчивый на holdout 2025."
        ]
    )

    add_heading(doc, "10. Ответы на контрольные вопросы")
    add_numbered(
        doc,
        [
            "Временной ряд состоит из тренда, сезонной и случайной компоненты. Тренд отражает долгосрочное направление, сезонность — регулярные месячные колебания, случайная компонента — нерегулярные отклонения.",
            "Коэффициент сезонности показывает отклонение месяца от среднего уровня. Значение 1,25 для декабря означает, что показатель в декабре в среднем на 25% выше обычного уровня.",
            "В аддитивной модели сезонность прибавляется к уровню ряда, в мультипликативной — умножается на него. Если амплитуда колебаний растет вместе с уровнем ряда, лучше использовать мультипликативную модель.",
            "MAPE — средняя абсолютная процентная ошибка. Значения ниже 10% обычно считаются высокой точностью, но при малых значениях ряда MAPE может быть неустойчивой.",
            "Прогноз инфляции можно использовать для бюджетирования, индексации цен, планирования закупок и оценки будущих расходов."
        ]
    )

    doc.save(LR3_PATH)


def build_lr4():
    raw, df, metrics, predictions, seasonality, future, features = load_data()

    doc = setup_doc()
    add_title_page(
        doc,
        "4",
        "«Применение методов машинного обучения для прогнозирования инфляции»"
    )

    add_heading(doc, "1. Цель работы и постановка бизнес-задачи")
    add_paragraph(
        doc,
        "Цель работы — применить методы машинного обучения для решения профильной бизнес-задачи. "
        "В данной работе рассматривается задача регрессии: прогноз месячного темпа инфляции "
        "по календарным, категориальным и временным признакам."
    )
    add_paragraph(
        doc,
        "Бизнес-смысл задачи заключается в поддержке планирования цен, бюджетирования, индексации расходов "
        "и оценки будущей динамики потребительских цен."
    )

    add_heading(doc, "2. Исследовательский анализ данных")
    add_table_caption(doc, "Таблица 1 — Паспорт набора данных")
    add_table(
        doc,
        ["№", "Характеристика", "Значение / описание"],
        [
            [1, "Название датасета", "Индексы потребительских цен по РФ"],
            [2, "Источник", "data/inflation_long_format.csv; первичный источник — Росстат"],
            [3, "Количество строк raw", len(raw)],
            [4, "Количество строк после подготовки", len(df)],
            [5, "Количество признаков модели", len(features["feature_cols"])],
            [6, "Тип задачи", "Регрессия"],
            [7, "Целевая переменная", "target = idx - 100"],
            [8, "Категориальные признаки", "Месяц и тип инфляции, закодированные one-hot"],
            [9, "Пропуски", f"{int(raw.isna().sum().sum())} пропусков"],
            [10, "Выбросы", "Высокие значения начала 1990-х; оставлены как часть экономической истории ряда"]
        ],
        widths=[1.5, 5, 8.5]
    )
    add_picture(
        doc,
        "target_distribution.png",
        "Рисунок 1 — Распределение целевой переменной target"
    )
    add_picture(
        doc,
        "time_series.png",
        "Рисунок 2 — Динамика месячной инфляции по категориям"
    )
    add_paragraph(
        doc,
        "EDA показывает, что в данных присутствует сильная асимметрия из-за периода высокой инфляции "
        "в начале 1990-х. Современный период имеет гораздо меньший разброс, поэтому абсолютные ошибки "
        "MAE и RMSE интерпретируются надежнее, чем один только R2."
    )

    add_heading(doc, "3. Предобработка данных")
    add_table_caption(doc, "Таблица 2 — Журнал предобработки данных")
    add_table(
        doc,
        ["Этап", "Операция", "Обоснование"],
        [
            ["Дата", "Преобразование в datetime", "Корректная сортировка и выделение года/месяца"],
            ["Категории", "One-Hot Encoding для inflation_type", "Модель получает различия между food, non_food, services и total"],
            ["Месяц", "One-Hot Encoding m_01 ... m_12", "Учет сезонной компоненты"],
            ["Лаги", "lag_1, lag_2, lag_3, lag_6, lag_12", "Использование исторической динамики ряда"],
            ["Скользящие средние", "roll_3, roll_6, roll_12", "Сглаживание краткосрочного шума"],
            ["Масштабирование", "StandardScaler для числовых признаков", "Необходимо для линейных и робастных моделей"],
            ["Разбиение", "Train/validation/test/holdout по времени", "Исключает перемешивание будущего с прошлым"]
        ],
        widths=[3, 5, 7]
    )

    add_heading(doc, "4. Построение и обучение моделей")
    add_paragraph(
        doc,
        "В работе обучены четыре модели scikit-learn. Для Ridge, HuberRegressor и GradientBoostingRegressor "
        "использован подбор гиперпараметров через GridSearchCV с TimeSeriesSplit."
    )
    add_table_caption(doc, "Таблица 3 — Использованные модели машинного обучения")
    add_table(
        doc,
        ["Модель", "Назначение", "Особенность"],
        [
            ["Linear Regression", "Базовая линейная регрессия", "Простой интерпретируемый baseline"],
            ["Ridge", "Регуляризованная линейная модель", "Снижает переобучение коэффициентов"],
            ["HuberRegressor", "Робастная регрессия", "Менее чувствительна к выбросам"],
            ["GradientBoostingRegressor", "Ансамблевая нелинейная модель", "Лучше улавливает нелинейные зависимости"]
        ],
        widths=[4.5, 5, 5.5]
    )

    add_heading(doc, "5. Оценка качества моделей")
    for table_number, (dataset_name, title) in enumerate([
        ("val_2019_2022", "Validation 2019–2022"),
        ("test_2023_2024", "Test 2023–2024"),
        ("holdout_2025", "Holdout 2025")
    ], start=4):
        subset = metrics[metrics["dataset"] == dataset_name].sort_values("mae")
        add_paragraph(doc, f"Период оценки: {title}.", bold_prefix="Период оценки:")
        add_table_caption(doc, f"Таблица {table_number} — Сравнение качества моделей: {title}")
        add_table(
            doc,
            ["Модель", "MAE", "MSE", "RMSE", "R2", "MAPE, %"],
            [
                [
                    row["model"],
                    fmt(row["mae"]),
                    fmt(row["mse"]),
                    fmt(row["rmse"]),
                    fmt(row["r2"]),
                    fmt(row["mape_idx"])
                ]
                for _, row in subset.iterrows()
            ],
            widths=[4.2, 2.1, 2.1, 2.1, 2.1, 2.4]
        )

    add_picture(
        doc,
        "test_mae.png",
        "Рисунок 3 — Сравнение моделей на тестовом периоде по MAE"
    )
    add_paragraph(
        doc,
        "Лучшей моделью по validation и test является Gradient Boosting. На holdout 2025 HuberRegressor "
        "показал более низкую ошибку, что говорит о его устойчивости к изменениям распределения."
    )

    add_heading(doc, "6. Анализ важности признаков")
    add_picture(
        doc,
        "feature_importance.png",
        "Рисунок 4 — Top-N важных признаков модели Gradient Boosting"
    )
    add_paragraph(
        doc,
        "Наиболее важные признаки связаны с предыдущими значениями инфляции и скользящими средними. "
        "Это подтверждает, что для прогноза месячной инфляции критична временная инерция ряда."
    )

    add_heading(doc, "7. Бизнес-рекомендации")
    add_bullets(
        doc,
        [
            "Использовать прогноз инфляции как входной показатель для бюджетирования и индексации расходов.",
            "Сравнивать ML-прогноз с сезонными baseline-моделями, чтобы не принимать переусложненную модель без практического выигрыша.",
            "Для стрессовых периодов дополнительно учитывать робастный прогноз HuberRegressor.",
            "Регулярно переобучать модель при поступлении новых месячных данных, так как структура инфляции меняется."
        ]
    )

    add_heading(doc, "8. Заключение")
    add_paragraph(
        doc,
        "В ходе работы была поставлена ML-задача регрессии, выполнены EDA и предобработка данных, "
        "обучены несколько моделей, рассчитаны метрики качества и интерпретированы результаты. "
        "Ограничение модели состоит в том, что используются только исторические значения инфляции; "
        "для повышения качества можно добавить курс валют, ключевую ставку, денежную массу и показатели внешней торговли."
    )

    add_heading(doc, "9. Ответы на контрольные вопросы")
    add_numbered(
        doc,
        [
            "Регрессия прогнозирует непрерывное числовое значение, классификация — класс объекта. Прогноз инфляции относится к регрессии, а прогноз ухода клиента — к классификации.",
            "Категориальные признаки кодируются в числа, потому что ML-алгоритмы работают с числовыми матрицами. One-Hot Encoding подходит для категорий без порядка, Ordinal Encoding — для упорядоченных категорий.",
            "При дисбалансе классов Accuracy может быть завышенной. Для churn-задачи важнее Precision, Recall, F1-score и ROC-AUC.",
            "Feature Importance показывает вклад признаков в прогноз модели. В задаче оттока важными могут быть срок договора, число обращений в поддержку и размер платежа; бизнес может точечно удерживать клиентов риска.",
            "Рост R2 с 0,85 до 0,90 после добавления расстояния до метро означает, что новый признак объясняет дополнительную часть вариации цены. Аналогично R2 полезен в прогнозировании спроса, выручки или уровня затрат."
        ]
    )

    doc.save(LR4_PATH)


def main():
    REPORTS_DIR.mkdir(exist_ok=True)
    build_lr3()
    build_lr4()
    print(LR3_PATH.resolve())
    print(LR4_PATH.resolve())


if __name__ == "__main__":
    main()
